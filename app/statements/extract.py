"""File → text/rows. Handles long PDFs page-by-page (streaming), CSV, XLSX/XLS, TXT, images (OCR)."""
from __future__ import annotations

import csv
import io
import logging
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

from app.core.config import get_settings

log = logging.getLogger(__name__)
ProgressCb = Callable[[int, int], None]

PDF_EXT = {".pdf"}
SHEET_EXT = {".csv", ".tsv", ".xlsx", ".xlsm", ".xls"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
TEXT_EXT = {".txt", ".text"}
SUPPORTED = PDF_EXT | SHEET_EXT | IMG_EXT | TEXT_EXT


class ExtractionError(Exception):
    pass


@dataclass
class Extracted:
    kind: str
    pages: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)  # tabular files
    ocr_pages: int = 0


def _ocr_image(img) -> str:
    try:
        import pytesseract
    except ImportError as e:
        raise ExtractionError("OCR needed but pytesseract isn't installed (pip install pytesseract + install Tesseract).") from e
    try:
        return pytesseract.image_to_string(img, config="--psm 6")
    except Exception as e:  # tesseract binary missing
        raise ExtractionError(f"OCR failed - is the Tesseract binary installed? ({e})") from e


def extract_pdf(path: Path, password: str | None, progress: ProgressCb | None) -> Extracted:
    import pdfplumber

    s = get_settings()
    out = Extracted(kind="pdf")
    try:
        pdf = pdfplumber.open(str(path), password=password or None)
    except Exception as e:
        msg = str(e).lower()
        if "password" in msg or "encrypt" in msg:
            raise ExtractionError("This PDF is password-protected. Re-upload with the password.") from e
        raise ExtractionError(f"Could not open PDF: {e}") from e
    with pdf:
        total = len(pdf.pages)
        for i, page in enumerate(pdf.pages, start=1):
            try:
                text = page.extract_text(x_tolerance=1.5, y_tolerance=3) or ""
            except Exception:
                text = ""
            if len(text.strip()) < 25 and s.ocr_enabled:  # scanned page
                try:
                    img = page.to_image(resolution=s.ocr_dpi).original
                    text = _ocr_image(img)
                    out.ocr_pages += 1
                except ExtractionError:
                    if i == 1:
                        raise
                except Exception as e:
                    log.warning("OCR failed on page %s: %s", i, e)
            out.pages.append(text)
            page.flush_cache()  # keep memory flat on very long PDFs
            if progress:
                progress(i, total)
    if not any(p.strip() for p in out.pages):
        raise ExtractionError("No readable text found in the PDF (scanned? enable OCR / install Tesseract).")
    return out


def _sheet_rows_xlsx(path: Path) -> list[list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    rows: list[list[str]] = []
    for ws in wb.worksheets:
        for r in ws.iter_rows(values_only=True):
            rows.append(["" if c is None else (c.strftime("%Y-%m-%d") if hasattr(c, "strftime") else str(c)) for c in r])
    return rows


def _sheet_rows_xls(path: Path) -> list[list[str]]:
    try:
        import xlrd
    except ImportError as e:
        raise ExtractionError("Legacy .xls needs xlrd (pip install xlrd) - or save the file as .xlsx/.csv.") from e
    wb = xlrd.open_workbook(str(path))
    rows = []
    for ws in wb.sheets():
        for r in range(ws.nrows):
            row = []
            for c in range(ws.ncols):
                cell = ws.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    dt = xlrd.xldate_as_datetime(cell.value, wb.datemode)
                    row.append(dt.strftime("%Y-%m-%d"))
                else:
                    row.append(str(cell.value))
            rows.append(row)
    return rows


def _sheet_rows_csv(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)]


def extract_file(path: Path, filename: str, password: str | None = None, progress: ProgressCb | None = None) -> Extracted:
    ext = Path(filename).suffix.lower()
    if ext in PDF_EXT:
        return extract_pdf(path, password, progress)
    if ext in {".xlsx", ".xlsm"}:
        rows = _sheet_rows_xlsx(path)
    elif ext == ".xls":
        rows = _sheet_rows_xls(path)
    elif ext in {".csv", ".tsv"}:
        rows = _sheet_rows_csv(path)
    elif ext in TEXT_EXT:
        text = path.read_text(errors="ignore")
        return Extracted(kind="text", pages=[text])
    elif ext in IMG_EXT:
        from PIL import Image

        return Extracted(kind="image", pages=[_ocr_image(Image.open(path))], ocr_pages=1)
    else:
        raise ExtractionError(f"Unsupported file type '{ext}'. Use PDF, CSV, XLSX/XLS, TXT or an image.")
    if progress:
        progress(1, 1)
    # tabular files also get a text rendering (used by the LLM fallback)
    text = "\n".join(" | ".join(c for c in r if c) for r in rows)
    return Extracted(kind="table", pages=[text], rows=rows)
